"""Acceptance 9.5: every table cell in the .xlsx equals the HTML report; charts are native (no images).

Usage: python tests/check_xlsx.py out/MBR_EST_Aug69.html out/MBR_EST_Aug69.xlsx
"""
from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from mbr.xlsx_out import parse_blocks, to_value


def main(html_path: str, xlsx_path: str) -> int:
    blocks = parse_blocks(Path(html_path).read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheets = wb.worksheets[1:]  # [0] is the table of contents
    errors, compared = [], 0
    if len(sheets) != len(blocks):
        errors.append(f"{len(sheets)} data sheets vs {len(blocks)} HTML tables blocks")
    for b, ws in zip(blocks, sheets):
        if ws.cell(1, 1).value != b.heading:
            errors.append(f"{ws.title}: title {ws.cell(1, 1).value!r} != heading {b.heading!r}")
        r = 3
        for t in b.tables:
            taken: set = set()
            for i, row in enumerate(t):
                c = 1
                for cell in row.cells:
                    while (r + i, c) in taken:
                        c += 1
                    got = ws.cell(r + i, c).value
                    want = cell.text if row.header or c == 1 else to_value(cell.text)[0]
                    compared += 1
                    if isinstance(want, float):
                        if not isinstance(got, (int, float)) or abs(got - want) > 1e-9:
                            errors.append(f"{ws.title} R{r + i}C{c}: {got!r} != {want!r} ({cell.text!r})")
                    elif (got or "") != want:
                        errors.append(f"{ws.title} R{r + i}C{c}: {got!r} != {want!r}")
                    for dr in range(cell.rowspan):
                        for dc in range(cell.colspan):
                            taken.add((r + i + dr, c + dc))
                    c += cell.colspan
            r += len(t) + 2
    with zipfile.ZipFile(xlsx_path) as z:
        names = z.namelist()
    charts = [n for n in names if re.match(r"xl/charts/chart\d+\.xml", n)]
    media = [n for n in names if n.startswith("xl/media/")]
    html_charts = Path(html_path).read_text(encoding="utf-8").count("<svg")
    print(f"compared {compared} cells in {len(sheets)} sheets; native charts: {len(charts)} "
          f"(HTML has {html_charts} SVG figures); embedded images: {len(media)}")
    if media:
        errors.append(f"embedded images found: {media}")
    if not charts:
        errors.append("no native charts")
    for e in errors[:40]:
        print("FAIL", e)
    print("OK" if not errors else f"{len(errors)} problem(s)")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main(*sys.argv[1:3]))
