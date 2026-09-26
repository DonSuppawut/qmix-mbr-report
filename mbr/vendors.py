"""Sub-contractor cartage by vendor, from `รายงานต้นทุน <Mon>'26 All Area.xlsx` sheet Exp&REV (spec 3.4)."""
from __future__ import annotations

import calendar
import json
import re
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

import openpyxl

from .fmt import EN_MONTH_ABBR

SUP_GROUP_TRUCK_RENT = "07"
# Vendor tag -> text patterns. Order matters: first match wins. TMP counts as QMix trucks (60 THB/m³ contract).
VENDORS = [
    ("TMP", "QMix", ("ไทย แมค พรีแคชท์",)),
    ("Indear99", "Sub", ("อินเดียร์99", "อินเดียร์ 99")),
    ("CCP", "Sub", ("Sub-CCP", "รถโม่ CCP", "ผลิตภัณฑ์คอนกรีตชลบุรี", "รถโม่พร้อมคนขับ")),
    ("UMO", "Sub", ("UMO",)),
    ("F-Transport", "Sub", ("F-Transport",)),
    ("Act Forward", "Sub", ("แอ็คท ฟอร์เวิร์ค", "Act Forward")),
]
_VOL = re.compile(r"([\d,]+(?:\.\d+)?)\s*(?:m3|M3|m³|คิว)")
THAI_MONTHS = ("ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค.")
# "1-15มิ.ย.69", "16-31 ก.ค.69" (Thai month, Buddhist year) or "1-15 Aug'26", "16-30 Jun'26"
_RANGE_TH = re.compile(r"(\d{1,2})\s*-\s*(\d{1,2})\s*(" + "|".join(re.escape(m) for m in THAI_MONTHS) + r")\s*(\d{2})")
_RANGE_EN = re.compile(r"(\d{1,2})\s*-\s*(\d{1,2})\s*(" + "|".join(EN_MONTH_ABBR) + r")\s*'?\s*(\d{2})")


def cost_report_path(account: Path, year: int, month: int) -> Path:
    return account / f"รายงานต้นทุน {EN_MONTH_ABBR[month - 1]}'{year % 100:02d} All Area.xlsx"


def _vendor(text: str) -> tuple[str, str] | None:
    for name, kind, pats in VENDORS:
        if any(p in text for p in pats):
            return name, kind
    return None


def date_range(text: str) -> tuple[date, date] | None:
    """Service period written in a GL text, e.g. "1-15มิ.ย.69" or "16-31 Jul'26" -> (start, end)."""
    m = _RANGE_TH.search(text)
    if m:
        mon, yr = THAI_MONTHS.index(m.group(3)) + 1, 2500 + int(m.group(4)) - 543
    else:
        m = _RANGE_EN.search(text)
        if not m:
            return None
        mon, yr = EN_MONTH_ABBR.index(m.group(3)) + 1, 2000 + int(m.group(4))
    last = calendar.monthrange(yr, mon)[1]
    d1, d2 = int(m.group(1)), min(int(m.group(2)), last)
    if not 1 <= d1 <= d2:
        return None
    return date(yr, mon, d1), date(yr, mon, d2)


def _days(r: tuple[date, date]) -> set[date]:
    return {r[0] + timedelta(n) for n in range((r[1] - r[0]).days + 1)}


def dispatch_snapshot_path(root: Path, year: int, month: int) -> Path:
    return root / "data" / "snapshots" / f"{year}-{month:02d}" / "ccp_dispatch_daily.json"


def add_ccp_dispatch_volume(v: dict, snapshot: Path, year: int, month: int) -> dict:
    """CCP invoice lines of this month that carry no m³ in the text (e.g. "ค่าบริการเช่ารถโม่พร้อมคนขับ 1-15 ก.ค.69")
    get their volume from production dispatches of the CCP trucks over the same days (QMixMBR.md 4.5).
    Lines whose days are already covered by a line with m³ (e.g. the insurance line next to "341.25 คิว 1-15 Aug'26")
    add nothing. A partial overlap is ambiguous — stop and ask Don."""
    no_vol = v.pop("ccp_ranges_no_vol", [])
    with_vol = v.pop("ccp_ranges_with_vol", [])
    covered: set[date] = set().union(*(_days(r) for r in with_vol)) if with_vol else set()
    need: set[date] = set()
    for r in no_vol:
        days = _days(r)
        if days <= covered:
            continue
        if days & covered:
            raise ValueError(f"{year}-{month:02d}: CCP line without m³ for {r[0]}..{r[1]} partly overlaps a line "
                             "with m³ — cannot tell which volume it covers; stop and check with Don")
        need |= days
    v["ccp_dispatch_days"] = sorted(need)
    if not need:
        return v
    if not snapshot.exists():
        raise FileNotFoundError(
            f"{snapshot} missing — fetch via query_eastsales: east_production_dispatches, truck_code in "
            f"east_sub_trucks_ccp.truck_no, production_date {min(need)}..{max(need)}, sum actual_vol by production_date")
    snap = json.loads(snapshot.read_text(encoding="utf-8"))
    daily = snap["daily"]
    lo, hi = snap["from"], snap["to"]
    if min(need).isoformat() < lo or max(need).isoformat() > hi:
        raise ValueError(f"{snapshot.name} covers {lo}..{hi} but {min(need)}..{max(need)} is needed")
    extra = sum(float(daily.get(d.isoformat(), 0.0)) for d in need)
    v["volume"]["CCP"] = v["volume"].get("CCP", 0.0) + extra
    v["ccp_dispatch_volume"] = extra
    return v


def vendor_month(path: Path, year: int, month: int) -> dict:
    """Amount = every vendor-tagged EST row of Sup.Group 07 (accruals, reversals, true-ups included) so the
    total ties to TVC row 164; volume = only rows whose text names the current month (e.g. "Aug'26")."""
    tag = f"{EN_MONTH_ABBR[month - 1]}'{year % 100:02d}".lower()
    amount: dict[str, float] = defaultdict(float)
    volume: dict[str, float] = defaultdict(float)
    kinds: dict[str, str] = {}
    untagged = 0.0
    ccp_with: list[tuple[date, date]] = []
    ccp_no: list[tuple[date, date]] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Exp&REV"]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        ix = {h: header.index(h) for h in ("Amount", "Text", "AO", "Sup. Group", "Ref.Doc")}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if "EST" not in str(row[ix["AO"]] or "") or not str(row[ix["Sup. Group"]] or "").startswith(SUP_GROUP_TRUCK_RENT):
                continue
            text = str(row[ix["Text"]] or "")
            v = _vendor(text) or _vendor(str(row[ix["Ref.Doc"]] or ""))
            amt = float(row[ix["Amount"]] or 0)
            if v is None:
                untagged += amt
                continue
            name, kind = v
            kinds[name] = kind
            amount[name] += amt
            m = _VOL.search(text)
            if tag in text.lower().replace(" ", "") and m:
                volume[name] += float(m.group(1).replace(",", ""))
            if name == "CCP" and amt > 0:
                r = date_range(text)
                if r and (r[0].year, r[0].month) == (year, month):
                    (ccp_with if m else ccp_no).append(r)
    finally:
        wb.close()
    return {"amount": dict(amount), "volume": dict(volume), "kind": kinds, "untagged": untagged,
            "ccp_ranges_with_vol": ccp_with, "ccp_ranges_no_vol": ccp_no}
