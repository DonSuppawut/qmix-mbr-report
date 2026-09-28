"""Reader for `MM.26 Net Con  (Management).xlsx` — per Site Code x Mix rows for Director=EST (Section 3)."""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .fmt import EN_MONTH_ABBR

SEGMENTS = ("Mega", "Large", "Medium", "Small", "Tiny")  # permanent display order (Don)
SEGMENT_LABELS = {"Mega": "Mega (20,001+)", "Large": "Large (5,001-20,000)", "Medium": "Medium (1,001-5,000)",
                  "Small": "Small (201-1,000)", "Tiny": "Tiny (0-200 m³)"}
# Customers counted as Mega on every row, whatever the file's segment says (Don 2026-09-24: China State).
MEGA_OVERRIDE = ("ไชน่า สเตท",)
CCP_CODE = "7312095"

# 1-based columns of the monthly sheet (header row 3)
COL = {"plant": 1, "director": 3, "aao": 5, "site": 7, "site_name": 8, "cust_code": 9, "cust_name": 10, "project_vol": 11, "segment": 12,
       "list_price": 13, "net_price": 15, "qty": 16, "total_net": 23, "cartage": 24, "rebate_bt": 35,
       "rawmat": 38, "tvc": 40, "netcon": 41}
HEADERS = {3: "Director", 5: "AAO", 7: "หน่วยงาน", 9: "รหัสลูกค้า", 11: "ปริมาณโครงการ", 12: "Segment", 16: "SUM_QTY",
           35: "Rebate เพิ่ม B/Ton", 41: "Nen Con (THB)"}


class NetconLayoutError(Exception):
    pass


@dataclass
class Row:
    site: str
    site_name: str
    cust_code: str
    cust_name: str
    project_vol: float
    segment: str
    list_price: float
    qty: float
    total_net: float
    cartage: float
    rawmat: float
    tvc: float
    netcon: float
    rebate_bt: float
    plant: str = ""
    aao: str = ""   # EST1/1, EST1/2, EST2


def netcon_path(account: Path, year: int, month: int) -> Path:
    return account / f"{month:02d}.{year % 100:02d} Net Con  (Management).xlsx"


def _bucket_segment(project_vol: float) -> str:
    if project_vol <= 200:
        return "Tiny"
    if project_vol <= 1000:
        return "Small"
    if project_vol <= 5000:
        return "Medium"
    if project_vol <= 20000:
        return "Large"
    return "Mega"


def classify(raw_segment: str, project_vol: float, customer: str) -> str:
    if any(c in customer for c in MEGA_OVERRIDE):
        return "Mega"
    seg = raw_segment.strip()
    if seg in SEGMENTS:
        return seg  # rows with an explicit segment are kept as-is
    return _bucket_segment(project_vol)


def read_netcon(path: Path, year: int, month: int) -> list[Row]:
    sheet = f"{EN_MONTH_ABBR[month - 1]}'{year % 100:02d}"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise NetconLayoutError(f"{path.name}: sheet {sheet!r} not found (wrong month file?) — {wb.sheetnames}")
        ws = wb[sheet]
        header = next(ws.iter_rows(min_row=3, max_row=3, max_col=51, values_only=True))
        for c, name in HEADERS.items():
            if str(header[c - 1] or "").strip() != name:
                raise NetconLayoutError(f"{path.name}/{sheet} col {c}: {header[c - 1]!r}, expected {name!r}")
        rows = []
        for r in ws.iter_rows(min_row=4, max_col=51, values_only=True):
            if r[COL["director"] - 1] != "EST":
                continue
            g = lambda k: r[COL[k] - 1]
            f = lambda k: float(g(k) or 0)
            cust = str(g("cust_name") or "")
            rows.append(Row(
                site=str(g("site") or ""), site_name=str(g("site_name") or ""), cust_code=str(g("cust_code") or ""),
                cust_name=cust, project_vol=f("project_vol"),
                segment=classify(str(g("segment") or ""), f("project_vol"), cust),
                list_price=f("list_price"), qty=f("qty"), total_net=f("total_net"), cartage=f("cartage"),
                rawmat=f("rawmat"), tvc=f("tvc"), netcon=f("netcon"), rebate_bt=f("rebate_bt"),
                plant=str(g("plant") or "").strip(), aao=str(g("aao") or "").strip()))
    finally:
        wb.close()
    return rows


# Net Con files name area E2 "EST2" (TVC: "EST2/1").
AREA_AAO = {"E1.1": ("EST1/1",), "E1.2": ("EST1/2",), "E2": ("EST2", "EST2/1")}


def area_rows(rows: list[Row], area: str) -> list[Row]:
    return [r for r in rows if r.aao in AREA_AAO[area]]


def by_segment(rows: list[Row]) -> dict[str, dict]:
    out = {s: defaultdict(float) for s in SEGMENTS}
    sites = {s: set() for s in SEGMENTS}
    for r in rows:
        d = out[r.segment]
        d["qty"] += r.qty
        d["list"] += r.list_price * r.qty
        for k in ("total_net", "cartage", "rawmat", "tvc", "netcon"):
            d[k] += getattr(r, k)
        sites[r.segment].add(r.site)
    for s in SEGMENTS:
        out[s]["projects"] = len(sites[s])
    return out


def total(rows: list[Row]) -> dict:
    d = defaultdict(float)
    for r in rows:
        d["qty"] += r.qty
        for k in ("total_net", "netcon"):
            d[k] += getattr(r, k)
    return d


def concentration(rows: list[Row], n: int) -> float:
    per = defaultdict(float)
    for r in rows:
        per[r.cust_code] += r.qty
    tot = sum(per.values())
    return sum(sorted(per.values(), reverse=True)[:n]) / tot * 100 if tot else 0.0


def customers(rows: list[Row]) -> dict[str, dict]:
    per: dict[str, dict] = defaultdict(lambda: defaultdict(float))
    for r in rows:
        per[r.cust_code]["qty"] += r.qty
        per[r.cust_code]["netcon"] += r.netcon
    return per


def sites(rows: list[Row]) -> dict[str, float]:
    per: dict[str, float] = defaultdict(float)
    for r in rows:
        per[r.site] += r.qty
    return per


NC_BUCKETS = [("<150 บาท/m³", None, 150), ("151-250 บาท/m³", 150, 250), ("251-350 บาท/m³", 250, 350),
              (">350 บาท/m³", 350, None)]


def nc_bucket(r: Row) -> str:
    nc = r.netcon / r.qty if r.qty else 0.0
    for label, lo, hi in NC_BUCKETS:
        if (lo is None or nc > lo) and (hi is None or nc <= hi):
            return label
    return NC_BUCKETS[-1][0]
