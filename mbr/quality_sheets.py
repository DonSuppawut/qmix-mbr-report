"""Plant strength sheets ('ตารางสรุปการเก็บข้อมูล กำลังอัด', one Google Sheet per plant, one tab per month).

Tabs are found by name (MM.2026 / MM/2026 → exported as MM2026 / optional '(…)' suffix per strength class) —
never by Google's gviz `sheet=` parameter, which silently returns another tab when the name does not exist.
Rows are read by their column-A labels, so the extra rows some plants keep (e.g. ปลวกแดง silt %) come along.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .fmt import TH_MONTH_FULL

ROOT = Path(__file__).resolve().parents[1]
FIRST_SAMPLE_COL, LAST_SAMPLE_COL = 2, 32  # B..AF


class QualitySheetError(Exception):
    pass


@dataclass
class MonthSheet:
    plant: str
    tab: str
    year: int
    month: int
    target: float | None
    product: str
    samples: list[int]                                  # sample numbers (row 'ตัวอย่างชุดที่')
    rows: list[tuple[str, list]] = field(default_factory=list)       # (label, value per sample)
    summary: list[tuple[str, dict]] = field(default_factory=list)    # (label, {avg, pct, margin, sd, cpk})
    series: list[tuple[str, list]] = field(default_factory=list)     # chart lines, same order as QC chart


def load_config() -> dict[str, str]:
    return json.loads((ROOT / "data" / "config" / "quality_sheets.json").read_text(encoding="utf-8"))["sheets"]


def to_dict(ms: MonthSheet) -> dict:
    return {"plant": ms.plant, "tab": ms.tab, "year": ms.year, "month": ms.month, "target": ms.target,
            "product": ms.product, "samples": ms.samples, "rows": [[lab, vals] for lab, vals in ms.rows],
            "summary": [[lab, d] for lab, d in ms.summary], "series": [[name, vals] for name, vals in ms.series]}


def from_dict(d: dict) -> MonthSheet:
    return MonthSheet(plant=d["plant"], tab=d["tab"], year=d["year"], month=d["month"], target=d["target"],
                      product=d["product"], samples=d["samples"], rows=[tuple(r) for r in d["rows"]],
                      summary=[tuple(s) for s in d["summary"]], series=[tuple(s) for s in d["series"]])


def snapshot_path(report_year: int, report_month: int, plant: str) -> Path:
    """Small, git-committed export of the parsed month sheet(s) actually used by one report — keeps QC and
    production-staff names for reference without committing the source workbook (1-3MB, whole multi-year sheet)."""
    return ROOT / "data" / "snapshots" / f"{report_year}-{report_month:02d}" / "quality" / f"{plant}.json"


def workbook_path(year: int, month: int, plant: str, file_id: str, refresh: bool = False) -> Path:
    path = ROOT / "data" / "snapshots" / f"{year}-{month:02d}" / "quality" / f"{plant}.xlsx"
    if refresh or not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        url = f"https://docs.google.com/spreadsheets/d/{file_id}/export?format=xlsx"
        with urllib.request.urlopen(url, timeout=60) as r:
            if "spreadsheetml" not in r.headers.get("Content-Type", ""):
                raise QualitySheetError(f"{plant}: download did not return an xlsx (sharing changed?) — {url}")
            path.write_bytes(r.read())
    return path


def month_tabs(names: list[str], year: int, month: int) -> list[str]:
    want = f"{month:02d}{year}"
    return [n for n in names if re.sub(r"\D", "", re.sub(r"\(.*?\)", "", n)) == want]


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        if not v or v.startswith("#"):  # '#DIV/0!' etc. from empty samples
            return None
    if isinstance(v, dt.datetime):
        return v.day
    return v


def _label(v) -> str:
    return " ".join(str(v).split()) if isinstance(v, str) else ""


def read_month(path: Path, plant: str, year: int, month: int) -> list[MonthSheet]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for tab in month_tabs(wb.sheetnames, year, month):
        ws = wb[tab]
        # ฏ/ฎ are commonly swapped when typing กรกฎาคม — compare with them folded together.
        fold = lambda t: t.replace(" ", "").replace("ฏ", "ฎ")
        month_text = fold(_label(ws.cell(2, 8).value))
        if month_text and fold(TH_MONTH_FULL[month - 1]) not in month_text:
            raise QualitySheetError(f"{plant} tab {tab!r}: header month is {month_text!r}, expected {TH_MONTH_FULL[month - 1]}")
        grid = {r: [ws.cell(r, c).value for c in range(1, LAST_SAMPLE_COL + 1)] for r in range(1, ws.max_row + 1)}
        labels = {r: _label(v[0]) for r, v in grid.items()}
        find = lambda prefix, after=0: next((r for r in sorted(labels) if r > after and labels[r].startswith(prefix)), None)
        r_no, r_first = find("ตัวอย่างชุดที่"), find("วันที่เก็บ")
        r_s1 = find("Strength (")
        r_s2 = find("Strength (", r_s1 or 0)
        if not (r_no and r_first and r_s1 and r_s2):
            raise QualitySheetError(f"{plant} tab {tab!r}: layout not recognised (sample / Strength rows missing)")
        cols = [c for c in range(FIRST_SAMPLE_COL, LAST_SAMPLE_COL + 1)
                if any(_clean(grid[r][c - 1]) is not None for r in range(r_first, r_s2 + 1))]
        target = _clean(ws.cell(2, 18).value)
        ms = MonthSheet(plant=plant, tab=tab, year=year, month=month,
                        target=float(target) if isinstance(target, (int, float)) else None,
                        product=_label(ws.cell(2, 24).value),
                        samples=[_clean(grid[r_no][c - 1]) or (i + 1) for i, c in enumerate(cols)])
        # Unlabelled rows right above an 'SD1 (x Days)' row hold the individual cube results for that age.
        cube_age: dict[int, str] = {}
        for r in sorted(labels):
            m = re.match(r"SD1 \((\d+) Days\)", labels[r])
            if m:
                k, rr = 1, r - 1
                stack = []
                while (rr > 0 and len(stack) < 3  # a sample has up to 3 cubes, in the 3 rows above its SD row
                       and (not labels[rr] or re.fullmatch(r"[\d.]+", labels[rr])) and rr not in cube_age):
                    stack.append(rr)
                    rr -= 1
                for n, row in enumerate(reversed(stack), 1):
                    cube_age[row] = f"ก้อนที่ {n} ({m.group(1)} วัน)"
        for r in sorted(grid):
            if r <= r_no:
                continue
            lab = labels[r]
            vals = [_clean(grid[r][c - 1]) for c in cols]
            if lab.startswith("กำลังอัดเฉลี่ย"):
                g = [_clean(x) for x in grid[r]]
                ms.summary.append((lab, {"avg": g[1], "pct": g[3], "margin": g[7], "sd": g[11], "cpk": g[15]}))
                continue
            if r in cube_age:
                lab = cube_age[r]
            elif not lab or re.fullmatch(r"[\d.]+", lab):
                continue  # e.g. the margin constant 40 in column A
            if all(v is None for v in vals):
                continue
            ms.rows.append((lab, vals))
        strength = [(lab, vals) for lab, vals in ms.rows if lab.startswith("Strength (")][:2]
        lines = strength + [next(((lab, vals) for lab, vals in ms.rows if lab.startswith(p)), None)
                            for p in ("เป้าหมาย", "กำลังรับรอง", "Margin")]
        for item in lines:
            if item:
                ms.series.append((item[0].replace(" (เฉลี่ย)", ""),
                                  [v if isinstance(v, (int, float)) else None for v in item[1]]))
        out.append(ms)
    return out
