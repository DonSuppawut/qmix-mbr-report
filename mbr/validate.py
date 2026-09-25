"""Phase-1 validators. Each returns a list of (ok, name, detail)."""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

import openpyxl

from . import calc
from .tvc import ScopeData

GOLDEN_JUL69 = {"volume": 25781.5, "ebitda": 3770974, "npat": 1769465, "cartage": 6098942, "assign_total": 2625778}
EXBT_TRD_GROUP = "11.Concrete from CPAC (Exbt&TRD)"
Check = tuple[bool, str, str]


def golden_jul69(jul: ScopeData) -> list[Check]:
    c = jul["combined"]
    got = {**{k: c[k] for k in ("volume", "ebitda", "npat", "cartage")}, "assign_total": calc.assign_total(c)}
    return [(abs(got[k] - v) < 1.0, f"golden Jul69 {k}", f"{got[k]:,.2f} vs {v:,.1f}") for k, v in GOLDEN_JUL69.items()]


def cost_report_totals(path: Path) -> dict[str, float]:
    """EST rows of sheet Exp&REV: Type='Cartage' total and Type='VCCost' excluding Exbt&TRD."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    tot: dict[str, float] = defaultdict(float)
    try:
        ws = wb["Exp&REV"]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        idx = {name: header.index(name) for name in ("Amount", "Type", "AO", "Sup. Group")}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if "EST" not in str(row[idx["AO"]] or ""):
                continue
            amt = float(row[idx["Amount"]] or 0)
            typ = row[idx["Type"]]
            if typ == "Cartage":
                tot["cartage"] += amt
            elif typ == "VCCost" and str(row[idx["Sup. Group"]] or "") != EXBT_TRD_GROUP:
                tot["assign_total"] += amt
    finally:
        wb.close()
    return tot


def reconcile_cost_report(actual: ScopeData, cost_report: Path) -> list[Check]:
    c = actual["combined"]
    cr = cost_report_totals(cost_report)
    return [
        (abs(cr["cartage"] - c["cartage"]) < 1.0, "cost report Cartage = TVC row 172",
         f"{cr['cartage']:,.2f} vs {c['cartage']:,.2f}"),
        (abs(cr["assign_total"] - calc.assign_total(c)) < 1.0, "cost report VCCost excl Exbt&TRD = TVC rows 192+195",
         f"{cr['assign_total']:,.2f} vs {calc.assign_total(c):,.2f}"),
    ]


def plant_sum_vs_scope(actual: ScopeData, pnl_snapshot: Path, plants_snapshot: Path) -> list[Check]:
    rows = json.loads(pnl_snapshot.read_text(encoding="utf-8"))["rows"]
    ptype = {p["plant_code"]: p["plant_type"] for p in json.loads(plants_snapshot.read_text(encoding="utf-8"))["rows"]}
    unknown = sorted({r["plant_code"] for r in rows} - ptype.keys())
    checks: list[Check] = [(not unknown, "all P&L plants known in east_plants", ", ".join(unknown) or "ok")]
    fields = {"volume": "volume_m3", "ebitda": "ebitda_thb", "npat": "npat_thb", "cartage": "cartage_thb"}
    for scope, types in (("combined", {"Core", "PMT"}), ("core", {"Core"}), ("pmt", {"PMT"})):
        for key, col in fields.items():
            s = sum(r[col] for r in rows if ptype.get(r["plant_code"]) in types)
            t = actual[scope][key]
            checks.append((abs(s - t) < 1.0, f"plant sum {scope} {key}", f"{s:,.2f} vs TVC {t:,.2f}"))
        s = sum(r["assign_thb"] for r in rows if ptype.get(r["plant_code"]) in types)
        t = calc.assign_total(actual[scope])
        checks.append((abs(s - t) < 1.0, f"plant sum {scope} assign_total", f"{s:,.2f} vs TVC {t:,.2f}"))
    return checks


def residual_checks(actual: ScopeData, base: ScopeData, label: str) -> list[Check]:
    out: list[Check] = []
    for scope in calc.SCOPES:
        for lvl, v in calc.residuals(actual[scope], base[scope]).items():
            out.append((abs(v) < 0.01, f"{label} residual {scope} {lvl}", f"{v:.6f}"))
    return out
