"""Truck registry — Google Sheet 'รถโม่กิจการตะวันออก' (link-shared, read-only), one tab per month named YYYY-MM.

Downloaded as .xlsx into data/snapshots/<ym>/truck_registry.xlsx (kept, so re-runs use the same list;
--refresh-sheets re-downloads). Returns the same shape the TU section uses:
    {"cols": [...], "trucks": [[truck, plant, owner, status, movement, note], ...],
     "merged": {"Q576+Q584": ["Q576", "Q584"]}, "fetched": "..."}
A replacement pair is merged when the note says the pair counts as one truck ('นับรวมเป็น 1 คัน').
"""
from __future__ import annotations

import datetime as dt
import re
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
HEADERS = {"เบอร์รถ": "truck", "รหัสโรงงาน": "plant", "รถQMix/ผรม.": "owner", "สถานะ": "status",
           "Movement": "movement", "หมายเหตุ": "note"}
COLS = ["truck", "plant", "owner", "status", "movement", "note"]
MERGE_MARK = "นับรวมเป็น 1 คัน"


class RegistryError(Exception):
    pass


def _download(file_id: str, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    url = f"https://docs.google.com/spreadsheets/d/{file_id}/export?format=xlsx"
    with urllib.request.urlopen(url, timeout=60) as r:
        if "spreadsheetml" not in r.headers.get("Content-Type", ""):
            raise RegistryError(f"ดาวน์โหลดทะเบียนรถไม่ได้ (การแชร์เปลี่ยน?) — {url}")
        path.write_bytes(r.read())


def load(year: int, month: int, file_id: str, refresh: bool = False) -> dict:
    ym = f"{year}-{month:02d}"
    path = ROOT / "data" / "snapshots" / ym / "truck_registry.xlsx"
    if refresh or not path.exists():
        _download(file_id, path)
    wb = openpyxl.load_workbook(path, data_only=True)
    if ym not in wb.sheetnames:
        raise RegistryError(f"ไฟล์ทะเบียนรถยังไม่มีชีท {ym} — เพิ่มชีทเดือนนี้ก่อน (copy ชีทล่าสุด → ตั้งชื่อ {ym}) "
                            f"แล้วรันใหม่ด้วย --refresh-sheets")
    rows = [r for r in wb[ym].iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
    head = [str(v or "").strip() for v in rows[0]]
    idx = {}
    for th, key in HEADERS.items():
        if th not in head:
            raise RegistryError(f"ชีท {ym}: ไม่พบคอลัมน์ {th!r} (หัวตาราง {head})")
        idx[key] = head.index(th)
    trucks = [[str(r[idx[k]] or "").strip() for k in COLS] for r in rows[1:] if r[idx["truck"]]]
    names = {t[0] for t in trucks}
    merged, done = {}, set()
    for t in trucks:
        if MERGE_MARK not in t[5] or t[0] in done:
            continue
        partner = next((p for p in re.findall(r"[A-Z]{1,2}\d{2,3}", t[5]) if p in names and p != t[0]), None)
        if partner is None:
            raise RegistryError(f"ชีท {ym}: {t[0]} ระบุ '{MERGE_MARK}' แต่หาเลขรถคู่ในหมายเหตุไม่เจอ: {t[5]!r}")
        pair = [t[0], partner]
        merged["+".join(pair)] = pair
        done.update(pair)
    fetched = dt.datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d")
    return {"cols": COLS, "trucks": trucks, "merged": merged, "fetched": fetched, "sheet": ym,
            "sheets": [n for n in wb.sheetnames if re.fullmatch(r"\d{4}-\d{2}", n)]}
